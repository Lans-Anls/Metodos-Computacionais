"""Orquestração dos cálculos e geração dos artefatos da aplicação."""
from __future__ import annotations

import math
from datetime import datetime
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from core import (
    ajuste_exponencial,
    ajuste_linear,
    ajuste_quadratico,
    bisseccao,
    compilar_expressao,
    erro_absoluto,
    gregory_newton,
    lagrange,
    limite_superior_erro,
    newton_diferencas_divididas,
    newton_raphson,
    newton_sistema,
    plotar,
    secante_amortecida,
)

from .configuracao import METODOS

BASE_DIR = Path(__file__).resolve().parent.parent


def _numero(texto: str, nome: str) -> float:
    try:
        valor = float(texto.strip().replace(",", "."))
    except ValueError as exc:
        raise ValueError(f"{nome}: informe um número válido.") from exc
    if not math.isfinite(valor):
        raise ValueError(f"{nome}: o valor deve ser finito.")
    return valor


def _inteiro_positivo(texto: str, nome: str) -> int:
    valor = _numero(texto, nome)
    if not valor.is_integer() or valor <= 0:
        raise ValueError(f"{nome}: informe um número inteiro maior que zero.")
    return int(valor)


def _numero_opcional(texto: str, nome: str) -> float | None:
    return None if not texto.strip() else _numero(texto, nome)


def _lista_numeros(texto: str, nome: str) -> list[float]:
    partes = texto.split(";") if ";" in texto else texto.split()
    if not partes or any(not parte.strip() for parte in partes):
        raise ValueError(f"{nome}: informe valores separados por ponto e vírgula.")
    return [_numero(parte, f"{nome}, item {indice + 1}")
            for indice, parte in enumerate(partes)]


def _pasta_saida(categoria: str) -> tuple[Path, Path]:
    nomes = {
        "Linear": "MetodosLineares",
        "Não linear": "Metodo de newton",
        "Interpolação": "Interpolacao",
        "Ajuste de curvas (MMQ)": "AjusteCurvas",
    }
    raiz = BASE_DIR / nomes[categoria]
    excel = raiz / "Excel"
    graficos = raiz / "IMG" / "Graficos"
    excel.mkdir(parents=True, exist_ok=True)
    graficos.mkdir(parents=True, exist_ok=True)
    return excel, graficos


def _salvar_excel(titulo: str, equacao: str, criterio: str, colunas: list[str],
                  tabela: list[dict], resumo: str, caminho: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Iterações"
    ws.append([titulo])
    ws.append([equacao])
    ws.append([f"Critério de parada: {criterio}"])
    ws.append([resumo])
    ws.append([])
    ws.append(colunas)

    azul = PatternFill("solid", fgColor="1F4E78")
    for celula in ws[6]:
        celula.fill = azul
        celula.font = Font(color="FFFFFF", bold=True)
        celula.alignment = Alignment(horizontal="center")

    for linha in tabela:
        ws.append([linha.get(coluna) for coluna in colunas])
    for coluna in ws.columns:
        largura = min(max(len(str(celula.value or "")) for celula in coluna) + 2, 32)
        ws.column_dimensions[coluna[0].column_letter].width = largura
    ws.freeze_panes = "A7"
    ws.auto_filter.ref = f"A6:{ws.cell(6, len(colunas)).coordinate}"
    wb.save(caminho)


def _grafico_sistema(linhas, raiz: np.ndarray, variaveis: tuple[str, ...],
                     caminho: Path) -> None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    ks = [linha[0] for linha in linhas]
    valores = np.array([linha[1] for linha in linhas])
    erros = [linha[3] for linha in linhas]
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(10, 4.8))
    for indice, rotulo in enumerate(variaveis):
        ax1.plot(ks, valores[:, indice], marker="o", label=rotulo)
    ax1.scatter([ks[-1] + 1] * len(variaveis), raiz,
                marker="x", s=70, color="#C62828")
    ax1.set(title="Evolução das variáveis", xlabel="Iteração k", ylabel="Valor")
    ax1.legend()
    ax1.grid(alpha=0.25)
    ax2.semilogy(ks, erros, marker="o", color="#C62828")
    ax2.set(title="Convergência do passo", xlabel="Iteração k", ylabel="máx|d|")
    ax2.grid(alpha=0.25)
    fig.suptitle(f"Newton para sistema não linear {len(variaveis)}x{len(variaveis)}",
                 fontweight="bold")
    fig.tight_layout()
    fig.savefig(caminho, dpi=140)
    plt.close(fig)


def _grafico_dados(xs: list[float], ys: list[float], avaliar, titulo: str,
                   caminho: Path, x_destaque: float | None = None) -> None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    minimo, maximo = min(xs), max(xs)
    margem = max((maximo - minimo) * 0.1, 0.5)
    grade_x = np.linspace(minimo - margem, maximo + margem, 500)
    grade_y = np.array([avaliar(float(valor)) for valor in grade_x])
    fig, ax = plt.subplots(figsize=(9, 5.5))
    ax.scatter(xs, ys, color="#153243", s=55, zorder=4, label="Dados")
    ax.plot(grade_x, grade_y, color="#D95D39", lw=2, label="Modelo")
    if x_destaque is not None:
        y_destaque = avaliar(x_destaque)
        ax.scatter([x_destaque], [y_destaque], color="#2E7D32", marker="x",
                   s=90, zorder=5, label=f"Estimativa ({x_destaque:.5g}, {y_destaque:.5g})")
    ax.set(title=titulo, xlabel="x", ylabel="y")
    ax.grid(alpha=0.25)
    ax.legend()
    fig.tight_layout()
    fig.savefig(caminho, dpi=140)
    plt.close(fig)


def calcular(categoria: str, metodo: str, entradas: dict[str, str]):
    """Executa o método selecionado e devolve tabela, resumo e artefatos."""
    configuracao = METODOS[categoria][metodo]
    excel_dir, grafico_dir = _pasta_saida(categoria)
    identificador = datetime.now().strftime("%Y%m%d_%H%M%S")
    nome = metodo.lower().replace(" ", "_").replace("-", "_")
    caminho_excel = excel_dir / f"{nome}_{identificador}.xlsx"
    caminho_grafico = grafico_dir / f"{nome}_{identificador}.png"

    if categoria == "Interpolação":
        return _calcular_interpolacao(configuracao, metodo, entradas,
                                      caminho_excel, caminho_grafico)
    if categoria == "Ajuste de curvas (MMQ)":
        return _calcular_ajuste(configuracao, metodo, entradas,
                                caminho_excel, caminho_grafico)

    max_iter = _inteiro_positivo(entradas["max_iter"], "Máximo de iterações")

    if metodo == "Bissecção":
        equacao = entradas["equacao"].strip()
        funcao = compilar_expressao(equacao, ("x",))
        a = _numero(entradas["a"], "Limite inferior a")
        b = _numero(entradas["b"], "Limite superior b")
        tol = _numero(entradas["tol"], "Tolerância")
        if a >= b or tol <= 0:
            raise ValueError("Use a < b e uma tolerância maior que zero.")
        resultado = bisseccao(funcao, a, b, tol, max_iter)
        intervalo = (a, b)
    elif metodo == "Newton-Raphson":
        equacao = entradas["equacao"].strip()
        funcao = compilar_expressao(equacao, ("x",))
        derivada = compilar_expressao(entradas["derivada"], ("x",))
        x0 = _numero(entradas["x0"], "Aproximação inicial x₀")
        resultado = newton_raphson(funcao, derivada, x0, 0.0, max_iter)
        alcance = max(8.0, abs(x0) * 1.5, abs(resultado.raiz or x0) * 1.25)
        intervalo = (-alcance if min(x0, resultado.raiz or x0) < 0 else 0.0, alcance)
    elif metodo == "Secante":
        equacao = entradas["equacao"].strip()
        funcao = compilar_expressao(equacao, ("x",))
        x0 = _numero(entradas["x0"], "Valor inicial x₀")
        x1 = _numero(entradas["x1"], "Valor inicial x₁")
        tol_pct = _numero(entradas["tol_pct"], "Tolerância relativa")
        if x0 == x1 or tol_pct <= 0:
            raise ValueError("Use valores iniciais distintos e tolerância maior que zero.")
        resultado = secante_amortecida(funcao, x0, x1, tol_pct / 100.0, max_iter)
        pontos = (x0, x1, resultado.raiz if resultado.raiz is not None else x1)
        margem = max(max(pontos) - min(pontos), 1.0) * 0.15
        intervalo = (min(pontos) - margem, max(pontos) + margem)
    else:
        return _calcular_sistema(configuracao, metodo, entradas, max_iter,
                                 caminho_excel, caminho_grafico)

    tabela = resultado.tabela
    raiz_texto = "indeterminada" if resultado.raiz is None else f"{resultado.raiz:.10g}"
    if metodo == "Newton-Raphson":
        resumo = (f"Resultado após {resultado.iteracoes} iterações: {raiz_texto}. "
                  "Execução por quantidade fixa concluída.")
        if resultado.motivo and resultado.motivo != "máximo de iterações atingido":
            resumo += f" Interrupção: {resultado.motivo}."
    else:
        chave_erro = "erro" if metodo == "Bissecção" else "erro_rel_%"
        erro_final = tabela[-1].get(chave_erro) if tabela else None
        erro_desejado = tol if metodo == "Bissecção" else tol_pct
        unidade = "%" if metodo == "Secante" else ""
        erro_texto = "indisponível" if erro_final is None else f"{erro_final:.6g}{unidade}"
        resumo = (f"Raiz aproximada: {raiz_texto}. Convergiu: "
                  f"{'sim' if resultado.convergiu else 'não'}; {resultado.iteracoes} iterações. "
                  f"Erro final: {erro_texto}; erro desejado: < {erro_desejado:.6g}{unidade}.")
        if resultado.motivo:
            resumo += f" Motivo: {resultado.motivo}."
        if not resultado.convergiu and resultado.motivo == "máximo de iterações atingido":
            unidade_iteracao = "iteração" if max_iter == 1 else "iterações"
            resumo += (f" A tolerância não foi alcançada no limite de {max_iter} "
                       f"{unidade_iteracao}; aumente esse limite ou revise os valores iniciais.")
    plotar(funcao, intervalo, [resultado.raiz], titulo=f"{metodo} - f(x) = {equacao}",
           salvar=str(caminho_grafico))
    _salvar_excel(metodo, f"f(x) = {equacao}", configuracao["criterio"],
                  resultado.colunas, tabela, resumo, caminho_excel)
    return resultado.colunas, tabela, resumo, caminho_excel, caminho_grafico


def _calcular_interpolacao(configuracao, metodo: str, entradas: dict[str, str],
                           caminho_excel: Path, caminho_grafico: Path):
    xs = _lista_numeros(entradas["xs"], "Valores de x")
    ys = _lista_numeros(entradas["ys"], "Valores de y")
    x_interp = _numero(entradas["x_interp"], "Valor de x a interpolar")
    if metodo == "Lagrange":
        resultado = lagrange(xs, ys, x_interp)
    elif metodo == "Newton - Diferenças Divididas":
        resultado = newton_diferencas_divididas(xs, ys, x_interp)
    else:
        resultado = gregory_newton(xs, ys, x_interp)

    coeficientes = ", ".join(
        f"a{indice}={valor:.8g}" for indice, valor in enumerate(resultado.coeficientes))
    resumo = (f"P{resultado.grau}({x_interp:.8g}) = {resultado.valor:.10g}. "
              f"Coeficientes: {coeficientes}.")
    valor_exato = _numero_opcional(entradas.get("valor_exato", ""), "Valor exato")
    if valor_exato is not None:
        resumo += f" Erro absoluto: {erro_absoluto(valor_exato, resultado.valor):.8g}."
    max_derivada = _numero_opcional(
        entradas.get("max_derivada", ""), "Máximo da derivada")
    if max_derivada is not None:
        limite = limite_superior_erro(xs, x_interp, max_derivada)
        resumo += f" Limite superior informado: {limite:.8g}."

    colunas = list(resultado.tabela[0])
    descricao = (f"x=[{', '.join(f'{valor:g}' for valor in xs)}]; "
                 f"y=[{', '.join(f'{valor:g}' for valor in ys)}]; x={x_interp:g}")
    _grafico_dados(xs, ys, resultado.avaliar, metodo, caminho_grafico, x_interp)
    _salvar_excel(metodo, descricao, configuracao["criterio"],
                  colunas, resultado.tabela, resumo, caminho_excel)
    return colunas, resultado.tabela, resumo, caminho_excel, caminho_grafico


def _calcular_ajuste(configuracao, metodo: str, entradas: dict[str, str],
                     caminho_excel: Path, caminho_grafico: Path):
    xs = _lista_numeros(entradas["xs"], "Valores de x")
    ys = _lista_numeros(entradas["ys"], "Valores de y")
    if metodo == "Ajuste Linear":
        resultado = ajuste_linear(xs, ys)
    elif metodo == "Ajuste Quadrático":
        resultado = ajuste_quadratico(xs, ys)
    else:
        resultado = ajuste_exponencial(xs, ys)

    coeficientes = ", ".join(
        f"c{indice}={valor:.8g}" for indice, valor in enumerate(resultado.coeficientes))
    resumo = (f"Modelo {resultado.modelo}: {coeficientes}. "
              f"R²={resultado.r_quadrado:.8g}; "
              f"soma dos quadrados dos resíduos={resultado.soma_quadrados_residuos:.8g}.")
    x_estimativa = _numero_opcional(
        entradas.get("x_estimativa", ""), "x para estimar y")
    if x_estimativa is not None:
        resumo += f" y({x_estimativa:.8g})={resultado.avaliar(x_estimativa):.10g}."

    colunas = list(resultado.tabela[0])
    descricao = (f"x=[{', '.join(f'{valor:g}' for valor in xs)}]; "
                 f"y=[{', '.join(f'{valor:g}' for valor in ys)}]")
    _grafico_dados(xs, ys, resultado.avaliar, metodo, caminho_grafico, x_estimativa)
    _salvar_excel(metodo, descricao, configuracao["criterio"],
                  colunas, resultado.tabela, resumo, caminho_excel)
    return colunas, resultado.tabela, resumo, caminho_excel, caminho_grafico


def _calcular_sistema(configuracao, metodo: str, entradas: dict[str, str],
                      max_iter: int, caminho_excel: Path, caminho_grafico: Path):
    variaveis = configuracao["variaveis"]
    equacoes = [entradas[chave].strip() for chave in configuracao["equacoes"]]
    funcoes = [compilar_expressao(equacao, variaveis) for equacao in equacoes]
    equacao = "; ".join(f"f{indice}={texto}" for indice, texto in enumerate(equacoes, 1))
    chute = tuple(_numero(entradas[chave], f"Valor inicial {variavel}")
                  for chave, variavel in zip(configuracao["iniciais"], variaveis))
    tol = _numero(entradas["tol"], "Tolerância")
    if tol <= 0:
        raise ValueError("A tolerância deve ser maior que zero.")
    try:
        linhas, raiz = newton_sistema(funcoes, chute, max_iter, tol)
    except np.linalg.LinAlgError as exc:
        raise ValueError("A Jacobiana ficou singular. Escolha outro passo inicial.") from exc

    tabela = []
    for k, ponto, valores, erro, passo in linhas:
        proximo = ponto - passo
        linha = {"k": k}
        linha.update({variavel: ponto[indice] for indice, variavel in enumerate(variaveis)})
        linha.update({f"f{indice + 1}": valores[indice] for indice in range(len(variaveis))})
        linha.update({f"d{indice + 1}": passo[indice] for indice in range(len(variaveis))})
        linha.update({f"{variavel}(k+1)": proximo[indice]
                      for indice, variavel in enumerate(variaveis)})
        linha["erro"] = erro
        tabela.append(linha)

    colunas = list(tabela[0])
    convergiu = bool(linhas and linhas[-1][3] < tol)
    erro_final = linhas[-1][3]
    raiz_formatada = ", ".join(f"{valor:.8f}" for valor in raiz)
    resumo = (f"Raiz aproximada: ({raiz_formatada}). "
              f"Convergiu: {'sim' if convergiu else 'não'}; {len(linhas)} iterações. "
              f"Erro final: {erro_final:.6g}; erro desejado: < {tol:.6g}.")
    if not convergiu:
        unidade_iteracao = "iteração" if max_iter == 1 else "iterações"
        resumo += (f" A tolerância não foi alcançada no limite de {max_iter} "
                   f"{unidade_iteracao}; aumente esse limite ou revise o ponto inicial.")
    _grafico_sistema(linhas, raiz, variaveis, caminho_grafico)
    _salvar_excel(metodo, equacao, configuracao["criterio"],
                  colunas, tabela, resumo, caminho_excel)
    return colunas, tabela, resumo, caminho_excel, caminho_grafico
